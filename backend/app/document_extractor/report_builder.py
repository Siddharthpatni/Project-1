"""
Report builder: converts TenderFields + source documents into a PDF or DOCX.

Structure of every report:
  1. Header (branding + metadata)
  2. Executive Summary — paragraph + bullet points (deterministic, no LLM)
  3. Structured Fields Table — all extracted key-value pairs
  4. Award Criteria / Eligibility Criteria
  5. Lots
  6. Full-text excerpt (truncated)
  7. Footer

No LLM — purely structural document generation using reportlab (PDF)
and python-docx (DOCX).
"""
from __future__ import annotations

import io
from datetime import datetime, UTC
from typing import Literal

from app.document_extractor.field_extractor import TenderFields
from app.document_extractor.parsers import ParsedDocument


# ---------------------------------------------------------------------------
# PDF builder (reportlab)
# ---------------------------------------------------------------------------

def build_pdf(
    fields: TenderFields,
    source_docs: list[ParsedDocument],
    source_url: str = "",
) -> bytes:
    try:
        from reportlab.lib import colors  # type: ignore
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import (
            HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )
    except ImportError:
        raise RuntimeError("reportlab is required: pip install reportlab") from None

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        title="Vergabepilot.AI — Extrahierte Ausschreibungsdaten",
    )
    styles = getSampleStyleSheet()

    BRAND  = colors.HexColor("#1a3c5e")
    ACCENT = colors.HexColor("#3b7dd8")
    LIGHT  = colors.HexColor("#f0f4fa")
    MUTED  = colors.HexColor("#6b7a8d")
    BULLET_BG = colors.HexColor("#eef3fb")

    h1_style = ParagraphStyle("H1", parent=styles["Heading1"],
                               textColor=BRAND, fontSize=16, spaceAfter=6)
    h2_style = ParagraphStyle("H2", parent=styles["Heading2"],
                               textColor=BRAND, fontSize=12, spaceBefore=14, spaceAfter=4)
    body_style = ParagraphStyle("Body", parent=styles["Normal"],
                                 fontSize=9, leading=13)
    label_style = ParagraphStyle("Label", parent=styles["Normal"],
                                  fontSize=8, textColor=MUTED, leading=11)
    summary_style = ParagraphStyle("Summary", parent=styles["Normal"],
                                    fontSize=10, leading=15, textColor=colors.HexColor("#1e293b"),
                                    spaceBefore=4, spaceAfter=4)
    bullet_style = ParagraphStyle("Bullet", parent=styles["Normal"],
                                   fontSize=9, leading=13, leftIndent=12,
                                   textColor=colors.HexColor("#1e293b"))
    caption_style = ParagraphStyle("Caption", parent=styles["Normal"],
                                    fontSize=7, textColor=MUTED, leading=10)

    story = []

    # ── Header ──────────────────────────────────────────────────────────────
    story.append(Paragraph("Vergabepilot.AI", h1_style))
    story.append(Paragraph("Automatisch extrahierte Ausschreibungsdaten", body_style))
    story.append(Spacer(1, 0.3 * cm))
    story.append(HRFlowable(width="100%", thickness=2, color=ACCENT))
    story.append(Spacer(1, 0.4 * cm))

    meta_rows = [
        ["Extrahiert am", datetime.now(UTC).strftime("%d.%m.%Y %H:%M UTC")],
        ["Quell-URL", source_url or "—"],
        ["Quelldokumente", ", ".join(d.filename for d in source_docs) or "—"],
    ]
    meta_tbl = Table(meta_rows, colWidths=[4 * cm, 13 * cm])
    meta_tbl.setStyle(TableStyle([
        ("FONTSIZE",   (0, 0), (-1, -1), 8),
        ("TEXTCOLOR",  (0, 0), (0, -1), MUTED),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING",    (0, 0), (-1, -1), 3),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 0.6 * cm))

    # ── Executive Summary ────────────────────────────────────────────────────
    story.append(Paragraph("Kurzzusammenfassung", h2_style))

    if fields.zusammenfassung:
        # Summary paragraph in a shaded box
        summary_box = Table(
            [[Paragraph(fields.zusammenfassung, summary_style)]],
            colWidths=[17 * cm],
        )
        summary_box.setStyle(TableStyle([
            ("BACKGROUND",    (0, 0), (-1, -1), BULLET_BG),
            ("ROUNDEDCORNERS", (0, 0), (-1, -1), [4, 4, 4, 4]),
            ("TOPPADDING",    (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING",   (0, 0), (-1, -1), 10),
            ("RIGHTPADDING",  (0, 0), (-1, -1), 10),
            ("BOX",           (0, 0), (-1, -1), 0.5, ACCENT),
        ]))
        story.append(summary_box)
        story.append(Spacer(1, 0.3 * cm))

    if fields.kernpunkte:
        story.append(Paragraph("Kernpunkte auf einen Blick:", h2_style))
        for point in fields.kernpunkte:
            story.append(Paragraph(f"&#x2022; &nbsp;{point}", bullet_style))
        story.append(Spacer(1, 0.4 * cm))

    # ── Structured Fields Table ──────────────────────────────────────────────
    story.append(Paragraph("Kerndaten der Ausschreibung", h2_style))

    field_rows: list[tuple[str, str]] = [
        ("Vergabenummer",    fields.vergabenummer or "—"),
        ("TED-Referenz",     fields.ted_reference or "—"),
        ("Auftraggeber",     fields.auftraggeber or "—"),
        ("Vergabestelle",    fields.vergabestelle or fields.auftraggeber or "—"),
        ("Titel",            fields.titel or "—"),
        ("Leistungsbeschreibung", (fields.leistungsbeschreibung or "—")[:300]),
        ("Vergabeverfahren", fields.vergabeverfahren or "—"),
        ("Auftragsart",      fields.auftragsart or "—"),
        ("Veröffentlicht",   fields.veroeffentlichungsdatum or "—"),
        ("Abgabefrist",      fields.abgabefrist or "—"),
        ("Bindefrist",       fields.bindefrist or "—"),
        ("CPV-Code(s)",      ", ".join(fields.cpv_codes) if fields.cpv_codes else "—"),
        ("NUTS-Code(s)",     ", ".join(fields.nuts_codes) if fields.nuts_codes else "—"),
        ("Auftragswert",     f"{fields.auftragswert} {fields.waehrung or ''}".strip() if fields.auftragswert else "—"),
        ("Leistungsort",     fields.leistungsort or "—"),
        ("Laufzeit",         fields.laufzeit or "—"),
        ("Ansprechpartner",  fields.ansprechpartner or "—"),
        ("E-Mail",           fields.email or "—"),
        ("Telefon",          fields.telefon or "—"),
        ("Fax",              fields.fax or "—"),
    ]
    # Only keep rows where the value is not "—"
    field_rows = [(label, val) for label, val in field_rows if val != "—"]

    tbl_data = [[Paragraph(f"<b>{label}</b>", label_style), Paragraph(val, body_style)]
                for label, val in field_rows]
    tbl = Table(tbl_data, colWidths=[5 * cm, 12 * cm])
    tbl.setStyle(TableStyle([
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT]),
        ("GRID",           (0, 0), (-1, -1), 0.3, colors.HexColor("#d0d8e8")),
        ("VALIGN",         (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING",     (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING",  (0, 0), (-1, -1), 5),
        ("LEFTPADDING",    (0, 0), (-1, -1), 6),
        ("RIGHTPADDING",   (0, 0), (-1, -1), 6),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 0.5 * cm))

    # ── Zuschlagskriterien ───────────────────────────────────────────────────
    if fields.zuschlagskriterien:
        story.append(Paragraph("Zuschlagskriterien", h2_style))
        for crit in fields.zuschlagskriterien:
            story.append(Paragraph(f"• {crit}", body_style))
        story.append(Spacer(1, 0.3 * cm))

    # ── Eignungskriterien ────────────────────────────────────────────────────
    if fields.eignungskriterien:
        story.append(Paragraph("Eignungskriterien", h2_style))
        for crit in fields.eignungskriterien:
            story.append(Paragraph(f"• {crit}", body_style))
        story.append(Spacer(1, 0.3 * cm))

    # ── Lose ─────────────────────────────────────────────────────────────────
    if fields.lose:
        story.append(Paragraph("Lose / Teillose", h2_style))
        for lot in fields.lose:
            story.append(Paragraph(f"• {lot}", body_style))
        story.append(Spacer(1, 0.3 * cm))

    # ── Full text excerpts ───────────────────────────────────────────────────
    for parsed in source_docs:
        if not parsed.text.strip():
            continue
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#d0d8e8")))
        story.append(Spacer(1, 0.3 * cm))
        story.append(Paragraph(f"Volltext: {parsed.filename}", h2_style))
        excerpt = parsed.text[:3000].replace("\n", "<br/>")
        if len(parsed.text) > 3000:
            excerpt += "<br/>...[gekürzt]"
        story.append(Paragraph(excerpt, body_style))
        story.append(Spacer(1, 0.3 * cm))

    # ── Footer ───────────────────────────────────────────────────────────────
    story.append(Spacer(1, 0.5 * cm))
    story.append(HRFlowable(width="100%", thickness=1, color=ACCENT))
    story.append(Spacer(1, 0.2 * cm))
    story.append(Paragraph(
        "Erstellt durch Vergabepilot.AI — Automatische Dokumentenanalyse · Deterministische Zusammenfassung · Keine KI",
        caption_style,
    ))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# DOCX builder (python-docx)
# ---------------------------------------------------------------------------

def build_docx(
    fields: TenderFields,
    source_docs: list[ParsedDocument],
    source_url: str = "",
) -> bytes:
    try:
        from docx import Document  # type: ignore
        from docx.shared import Pt, RGBColor, Cm  # type: ignore
    except ImportError:
        raise RuntimeError("python-docx is required: pip install python-docx") from None

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    BRAND = RGBColor(0x1a, 0x3c, 0x5e)

    # ── Title ────────────────────────────────────────────────────────────────
    title_p = doc.add_heading("Vergabepilot.AI — Ausschreibungsdaten", level=0)
    if title_p.runs:
        title_p.runs[0].font.color.rgb = BRAND

    doc.add_paragraph(f"Extrahiert: {datetime.now(UTC).strftime('%d.%m.%Y %H:%M UTC')}")
    if source_url:
        doc.add_paragraph(f"Quelle: {source_url}")
    doc.add_paragraph(f"Dokumente: {', '.join(d.filename for d in source_docs) or '—'}")

    # ── Executive Summary ────────────────────────────────────────────────────
    doc.add_heading("Kurzzusammenfassung", level=1)

    if fields.zusammenfassung:
        summary_para = doc.add_paragraph(fields.zusammenfassung)
        summary_para.style.font.size = Pt(10)

    if fields.kernpunkte:
        doc.add_heading("Kernpunkte auf einen Blick", level=2)
        for point in fields.kernpunkte:
            doc.add_paragraph(point, style="List Bullet")

    doc.add_paragraph("")

    # ── Structured fields table ──────────────────────────────────────────────
    doc.add_heading("Kerndaten der Ausschreibung", level=1)

    field_rows: list[tuple[str, str]] = [
        ("Vergabenummer",    fields.vergabenummer or "—"),
        ("TED-Referenz",     fields.ted_reference or "—"),
        ("Auftraggeber",     fields.auftraggeber or "—"),
        ("Vergabestelle",    fields.vergabestelle or fields.auftraggeber or "—"),
        ("Titel",            fields.titel or "—"),
        ("Leistungsbeschreibung", (fields.leistungsbeschreibung or "—")[:300]),
        ("Vergabeverfahren", fields.vergabeverfahren or "—"),
        ("Auftragsart",      fields.auftragsart or "—"),
        ("Veröffentlicht",   fields.veroeffentlichungsdatum or "—"),
        ("Abgabefrist",      fields.abgabefrist or "—"),
        ("Bindefrist",       fields.bindefrist or "—"),
        ("CPV-Code(s)",      ", ".join(fields.cpv_codes) if fields.cpv_codes else "—"),
        ("NUTS-Code(s)",     ", ".join(fields.nuts_codes) if fields.nuts_codes else "—"),
        ("Auftragswert",     f"{fields.auftragswert} {fields.waehrung or ''}".strip() if fields.auftragswert else "—"),
        ("Leistungsort",     fields.leistungsort or "—"),
        ("Laufzeit",         fields.laufzeit or "—"),
        ("Ansprechpartner",  fields.ansprechpartner or "—"),
        ("E-Mail",           fields.email or "—"),
        ("Telefon",          fields.telefon or "—"),
        ("Fax",              fields.fax or "—"),
    ]
    # Only keep rows with actual values
    field_rows = [(label, val) for label, val in field_rows if val != "—"]

    if field_rows:
        tbl = doc.add_table(rows=1, cols=2)
        tbl.style = "Table Grid"
        hdr = tbl.rows[0].cells
        hdr[0].text = "Feld"
        hdr[1].text = "Wert"
        for cell in hdr:
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for label, val in field_rows:
            row = tbl.add_row().cells
            row[0].text = label
            row[1].text = val

    doc.add_paragraph("")

    # ── Zuschlagskriterien ───────────────────────────────────────────────────
    if fields.zuschlagskriterien:
        doc.add_heading("Zuschlagskriterien", level=2)
        for crit in fields.zuschlagskriterien:
            doc.add_paragraph(crit, style="List Bullet")

    # ── Eignungskriterien ────────────────────────────────────────────────────
    if fields.eignungskriterien:
        doc.add_heading("Eignungskriterien", level=2)
        for crit in fields.eignungskriterien:
            doc.add_paragraph(crit, style="List Bullet")

    # ── Lose ─────────────────────────────────────────────────────────────────
    if fields.lose:
        doc.add_heading("Lose / Teillose", level=2)
        for lot in fields.lose:
            doc.add_paragraph(lot, style="List Bullet")

    # ── Full text ────────────────────────────────────────────────────────────
    for parsed in source_docs:
        if not parsed.text.strip():
            continue
        doc.add_heading(f"Volltext: {parsed.filename}", level=2)
        excerpt = parsed.text[:3000]
        if len(parsed.text) > 3000:
            excerpt += "\n...[gekürzt]"
        doc.add_paragraph(excerpt)

    # ── Footer ───────────────────────────────────────────────────────────────
    footer_p = doc.add_paragraph(
        "\n— Erstellt durch Vergabepilot.AI — Deterministische Analyse · Keine KI · Kein API-Aufruf"
    )
    footer_p.runs[0].font.size = Pt(8)
    footer_p.runs[0].font.color.rgb = RGBColor(0x6b, 0x7a, 0x8d)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Unified entry point
# ---------------------------------------------------------------------------

def build_report(
    fields: TenderFields,
    source_docs: list[ParsedDocument],
    source_url: str = "",
    fmt: Literal["pdf", "docx"] = "pdf",
) -> tuple[bytes, str]:
    """Return (bytes, mime_type)."""
    if fmt == "docx":
        return (
            build_docx(fields, source_docs, source_url),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    return build_pdf(fields, source_docs, source_url), "application/pdf"
